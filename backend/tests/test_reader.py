"""Tests for reading files in parts (P1-E3-T2): PDFs by page, Markdown and Word by section, any file by
line range, every read capped with the totals and the line it stopped at. Uses a temporary database
and folders; the PDF is written by a small helper, so no PDF library is needed.
"""

import os
import shutil
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from app.core import config
from app.db import database
from app.services import catalog_service, reader_service


def _write(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def _write_pdf(path, pages):
    """Writes a PDF with one line of text per page, in Helvetica."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>"
            % len(objects)
        )
        kids.append(len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % kid for kid in kids), len(kids))
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    with open(path, "wb") as f:
        f.write(out)


class ReaderTestCase(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="reader_test_")
        self.docs = os.path.join(self.tmpdir, "docs")
        os.makedirs(self.docs)
        self._patches = [
            patch.object(database, "DB_PATH", os.path.join(self.tmpdir, "assistant_test.db")),
            patch.object(database, "_initialized", False),
            # The temp folder may sit under AppData, which is never indexed; tests set their own rules.
            patch.object(config, "SYSTEM_EXCLUDE", [os.path.join(self.tmpdir, "Windows")]),
            patch.object(config, "PRIVACY_EXCLUDE", ["Google/Chrome"]),
            patch.object(config, "IGNORE_DIRS", {".git", "node_modules"}),
        ]
        for p in self._patches:
            p.start()
        self.project_id = database.create_project("Thesis", folders=[(self.docs, "read")])["id"]

    def tearDown(self):
        for p in self._patches:
            p.stop()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _file_id(self, name):
        """Scans the project and returns the catalog id of the file called `name`."""
        catalog_service.catalog_project(database.get_project(self.project_id))
        return next(f["id"] for f in database.list_catalog_files(self.project_id) if f["name"] == name)

    def _read(self, name, **where):
        return reader_service.read_part(self.project_id, self._file_id(name), **where)


class AcceptanceTests(ReaderTestCase):
    def test_a_300_page_pdf_can_be_read_page_by_page(self):
        _write_pdf(os.path.join(self.docs, "report.pdf"), [f"Page {n} of the report" for n in range(1, 301)])
        file_id = self._file_id("report.pdf")

        for n in range(1, 301):
            part = reader_service.read_part(self.project_id, file_id, page=n)
            self.assertEqual((part["unit"], part["page"], part["text"]), ("page", n, f"Page {n} of the report"))
            self.assertEqual(part["total_pages"], 300)
            self.assertFalse(part["truncated"])


class UnitTests(ReaderTestCase):
    def test_markdown_is_read_by_heading(self):
        text = "Intro line\n# Methods\nWe measured.\n```\n# not a heading\n```\n## Results\nIt worked."
        _write(os.path.join(self.docs, "notes.md"), text)

        part = self._read("notes.md", section=2)

        self.assertEqual((part["title"], part["start_line"], part["end_line"]), ("Methods", 2, 6))
        self.assertEqual(part["text"], "# Methods\nWe measured.\n```\n# not a heading\n```")
        self.assertEqual(part["total_sections"], 3)

    def test_word_headings_become_sections(self):
        doc = Document()
        doc.add_heading("Report", 0)
        doc.add_paragraph("Summary.")
        doc.add_heading("Budget", 1)
        doc.add_paragraph("Costs went down.")
        doc.save(os.path.join(self.docs, "report.docx"))

        part = self._read("report.docx", section=2)

        self.assertEqual((part["title"], part["text"]), ("Budget", "# Budget\nCosts went down."))
        self.assertEqual(part["total_sections"], 2)

    def test_code_is_read_by_line_range(self):
        _write(os.path.join(self.docs, "main.py"), "\n".join(f"line {n}" for n in range(1, 6)))

        part = self._read("main.py", start_line=2, end_line=3)

        self.assertEqual((part["unit"], part["text"], part["start_line"], part["end_line"]), ("lines", "line 2\nline 3", 2, 3))
        self.assertEqual((part["total_lines"], part["total_sections"]), (5, None))


class CapTests(ReaderTestCase):
    def test_capped_reads_continue_until_the_whole_file_is_read(self):
        text = "\n".join(["abcd"] * 10)
        _write(os.path.join(self.docs, "notes.txt"), text)
        file_id = self._file_id("notes.txt")

        pieces, start = [], 1
        with patch.object(config, "MAX_READ_CHARS", 12):
            while True:
                part = reader_service.read_part(self.project_id, file_id, start_line=start)
                pieces.append(part["text"])
                if not part["truncated"]:
                    break
                start = part["end_line"] + 1

        self.assertEqual(len(pieces), 5)
        self.assertEqual("\n".join(pieces), text)
        self.assertEqual(part["total_chars"], len(text))

    def test_a_line_longer_than_the_cap_is_cut(self):
        _write(os.path.join(self.docs, "bundle.js"), "x" * 50)

        with patch.object(config, "MAX_READ_CHARS", 20):
            part = self._read("bundle.js")

        self.assertEqual((len(part["text"]), part["truncated"]), (20, True))


class RefusalTests(ReaderTestCase):
    def test_a_page_that_doesnt_exist_is_refused(self):
        _write_pdf(os.path.join(self.docs, "short.pdf"), ["Only page"])

        with self.assertRaises(reader_service.PartNotFound):
            self._read("short.pdf", page=2)

    def test_each_file_is_read_only_its_own_way(self):
        _write_pdf(os.path.join(self.docs, "short.pdf"), ["Only page"])
        _write(os.path.join(self.docs, "main.py"), "print(1)")
        _write(os.path.join(self.docs, "notes.md"), "# One")

        for name, where in (("short.pdf", {"section": 1}), ("main.py", {"section": 1}), ("notes.md", {"page": 1})):
            with self.subTest(name=name), self.assertRaises(reader_service.PartNotFound):
                self._read(name, **where)

    def test_lines_outside_the_file_are_refused(self):
        _write(os.path.join(self.docs, "main.py"), "one\ntwo")

        with self.assertRaises(reader_service.PartNotFound):
            self._read("main.py", start_line=5)

    def test_a_file_not_in_the_catalog_returns_none(self):
        self.assertIsNone(reader_service.read_part(self.project_id, 999))


if __name__ == "__main__":
    unittest.main()
